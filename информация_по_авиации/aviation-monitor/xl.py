#!/usr/bin/env python3
"""Excel для Ави: понять файл, прочитать данные, проверить финмодель, вытащить отчётность по кодам строк.

  ~/pw/bin/python xl.py info  файл.xlsx                  листы, размеры, формулы, скрытые листы, графики, внешние ссылки
  ~/pw/bin/python xl.py dump  файл.xlsx [лист] [A1:H40]  значения таблицей (то, что видно в Excel)
  ~/pw/bin/python xl.py check файл.xlsx                  проверка модели: ошибки, «зашитые» числа, внешние ссылки
  ~/pw/bin/python xl.py codes файл.xlsx                  отчётность РСБУ по кодам строк (1600, 2110…) → коэффициенты
  ~/pw/bin/python xl.py charts файл.xlsx                 графики: название, тип и диапазоны данных каждой серии

Исходный файл никогда не перезаписывается (openpyxl теряет графики при сохранении) — результат только в новый файл.
Значения формул берутся из кэша Excel: если файл сохраняли не в Excel, у формул может не быть значений («н/д»).
"""
import re
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

ERRORS = ("#REF!", "#DIV/0!", "#N/A", "#VALUE!", "#NAME?", "#NUM!", "#NULL!")
CODES = {"1100": "Внеоборотные активы", "1150": "Основные средства", "1200": "Оборотные активы",
         "1230": "Дебиторская задолженность", "1250": "Денежные средства", "1300": "Капитал и резервы",
         "1400": "Долгосрочные обязательства", "1410": "Долгоср. заёмные средства", "1500": "Краткосрочные обязательства",
         "1510": "Краткоср. заёмные средства", "1520": "Кредиторская задолженность", "1600": "Баланс (активы)",
         "2110": "Выручка", "2120": "Себестоимость", "2200": "Прибыль от продаж", "2330": "Проценты к уплате",
         "2300": "Прибыль до налогообложения", "2400": "Чистая прибыль"}


def fmt(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:,.2f}".replace(",", " ").replace(".", ",")
    return str(v)


def cmd_info(path):
    wb = load_workbook(path, data_only=False)
    print(f"Файл: {Path(path).name}")
    for ws in wb.worksheets:
        formulas = sum(1 for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
        state = "" if ws.sheet_state == "visible" else f" ⚠️ лист {ws.sheet_state}"
        print(f"  лист «{ws.title}»: {ws.dimensions}, формул {formulas}{state}")
    if wb.defined_names:
        print(f"  именованных диапазонов: {len(list(wb.defined_names))}")
    ext = external_links(path)
    if ext:
        print(f"  ⚠️ внешние ссылки на другие файлы: {len(ext)} (значения могут быть устаревшими)")
    n = len(chart_files(path))
    print(f"  графиков: {n}" + (" — подробности: xl.py charts" if n else ""))


def cmd_dump(path, sheet=None, rng=None):
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    if rng:
        c1, r1, c2, r2 = range_boundaries(rng)
    else:
        c1, r1, c2, r2 = 1, 1, min(ws.max_column, 15), min(ws.max_row, 60)
    print(f"«{ws.title}» {rng or f'первые {r2} строк × {c2} столбцов'}")
    for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
        vals = [fmt(c.value) for c in row]
        if any(vals):
            print(f"{row[0].row:>4} | " + " | ".join(vals))


def cmd_check(path):
    """Типовые ошибки финмоделей: ошибки в ячейках, числа «зашиты» в формулы, внешние ссылки, скрытые листы."""
    wb_f = load_workbook(path, data_only=False)
    wb_v = load_workbook(path, data_only=True)
    issues = []
    for ws in wb_f.worksheets:
        vs = wb_v[ws.title]
        if ws.sheet_state != "visible":
            issues.append(f"лист «{ws.title}» скрыт — проверить, что в нём")
        for row in ws.iter_rows():
            for c in row:
                v, cached = c.value, vs[c.coordinate].value
                if isinstance(cached, str) and cached.strip() in ERRORS:
                    issues.append(f"{ws.title}!{c.coordinate}: ошибка {cached}")
                elif isinstance(v, str) and v.startswith("="):
                    if any(e in v for e in ERRORS):
                        issues.append(f"{ws.title}!{c.coordinate}: битая ссылка в формуле")
                    # «зашитое» число в формуле: =B5*1,2 или =C10+150000 — частая причина скрытых допущений
                    nums = [n for n in re.findall(r"(?<![A-Z$\d.])\d+(?:\.\d+)?(?![\d:!])", v.upper())
                            if n not in ("0", "1", "100", "12", "365")]
                    if nums:
                        issues.append(f"{ws.title}!{c.coordinate}: в формуле зашито число {', '.join(nums[:3])} → {v[:60]}")
                    if "[" in v and "]" in v:
                        issues.append(f"{ws.title}!{c.coordinate}: ссылка на другой файл → {v[:60]}")
            # число вбито руками посреди строки формул (слева и справа — формулы): расчёт «перебит» вручную
            cells = list(row)
            for i in range(1, len(cells) - 1):
                left, mid, right = cells[i - 1].value, cells[i].value, cells[i + 1].value
                if (isinstance(mid, (int, float)) and not isinstance(mid, bool)
                        and all(isinstance(x, str) and x.startswith("=") for x in (left, right))):
                    issues.append(f"{ws.title}!{cells[i].coordinate}: число {mid} вместо формулы "
                                  f"(соседние ячейки считаются формулами) — проверить, не перебит ли расчёт")
    if external_links(path):
        issues.append("в книге есть связи с внешними файлами — значения могли не обновиться")
    print(f"Проверка {Path(path).name}: замечаний {len(issues)}")
    for i in issues[:60]:
        print("  • " + i)
    if len(issues) > 60:
        print(f"  … ещё {len(issues) - 60}")


def cmd_codes(path):
    """Ищет строки с кодами РСБУ (4 цифры) и числа справа — по каждому периоду-столбцу; считает основные коэффициенты."""
    wb = load_workbook(path, data_only=True)
    found = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for i, c in enumerate(row):
                code = str(c.value).strip() if c.value is not None else ""
                if code in CODES and code not in found:
                    nums = [x.value for x in row[i + 1:] if isinstance(x.value, (int, float))]
                    if nums:
                        found[code] = nums[:3]
    if not found:
        sys.exit("Кодов строк РСБУ (1600, 2110…) в файле не нашёл — это не стандартная форма отчётности.")
    print("Код  | Показатель                     | значения по столбцам (как в файле)")
    for code in sorted(found):
        print(f"{code} | {CODES[code]:<30} | " + " | ".join(fmt(float(v)) for v in found[code]))

    def g(code, k=0):
        v = found.get(code)
        return float(v[k]) if v and len(v) > k else None

    def ratio(a, b):
        return f"{a / b:.2f}".replace(".", ",") if a is not None and b else "н/д"

    print("\nКоэффициенты (по первому столбцу значений):")
    debt = sum(x for x in (g("1410"), g("1510")) if x is not None) if (g("1410") or g("1510")) else None
    print(f"  Доля собственного капитала = 1300 / 1600 = {ratio(g('1300'), g('1600'))}")
    print(f"  Текущая ликвидность = 1200 / 1500 = {ratio(g('1200'), g('1500'))}")
    print(f"  Чистая рентабельность = 2400 / 2110 = {ratio(g('2400'), g('2110'))}")
    print(f"  Покрытие процентов = (2300 + 2330) / 2330 = "
          f"{ratio((g('2300') or 0) + (g('2330') or 0), g('2330')) if g('2300') is not None else 'н/д'}")
    print(f"  Заёмные средства / капитал = (1410 + 1510) / 1300 = {ratio(debt, g('1300'))}")
    print("Проверь, что первый столбец — последний период: порядок столбцов в формах бывает разный.")


def chart_files(path):
    with zipfile.ZipFile(path) as z:
        return sorted(n for n in z.namelist() if re.match(r"xl/charts/chart\d+\.xml$", n))


def external_links(path):
    with zipfile.ZipFile(path) as z:
        return [n for n in z.namelist() if n.startswith("xl/externalLinks/")]


def cmd_charts(path):
    """Графики читаем из XML книги: тип, название и диапазоны данных — по ним Ави берёт сами числа через dump."""
    files = chart_files(path)
    if not files:
        print("Графиков в файле нет.")
        return
    with zipfile.ZipFile(path) as z:
        for n in files:
            xml = z.read(n).decode("utf-8", "replace")
            # префикс пространства имён бывает «c:» или пустой — ищем без него
            kind = sorted(set(re.findall(r"<(?:\w+:)?(\w+Chart)\b", xml)))
            head = re.split(r"<(?:\w+:)?plotArea\b", xml)[0]
            title = " ".join(re.findall(r"<(?:\w+:)?t>([^<]*)</(?:\w+:)?t>", head))[:80]
            refs = re.findall(r"<(?:\w+:)?f>([^<]+)</(?:\w+:)?f>", xml)
            print(f"{n.split('/')[-1]}: {', '.join(kind) or 'тип не определён'} · «{title or 'без названия'}»")
            for r in refs[:12]:
                print(f"    данные: {r}")


if __name__ == "__main__":
    cmds = {"info": cmd_info, "dump": cmd_dump, "check": cmd_check, "codes": cmd_codes, "charts": cmd_charts}
    if len(sys.argv) < 3 or sys.argv[1] not in cmds:
        print(__doc__)
        sys.exit(1)
    cmds[sys.argv[1]](*sys.argv[2:])
