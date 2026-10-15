#!/usr/bin/env python3
"""Чтение PDF: отчётность, договоры, отраслевые обзоры. Только на сервере, наружу ничего не отправляет.

  ~/pw/bin/python pdfread.py info   файл.pdf                  страниц, есть ли текстовый слой (или это скан)
  ~/pw/bin/python pdfread.py text   файл.pdf [1-5]            текст по страницам → файл.txt рядом (и первые строки на экран)
  ~/pw/bin/python pdfread.py tables файл.pdf [1-5]            все таблицы → файл_таблицы.xlsx (лист на таблицу), числа — числами
  ~/pw/bin/python pdfread.py codes  файл.pdf                  строки РСБУ (1600, 2110…) из таблиц + коэффициенты (как xl.py codes)

Сканы (картинки без текста) не читаются: распознавание текста на этом сервере не установлено — так и сообщи.
Числа: «1 234 567» → 1234567; «(12 345)» и «−12 345» → отрицательные.
"""
import logging
import re
import sys
from pathlib import Path

import pdfplumber
from openpyxl import Workbook

logging.getLogger("pdfminer").setLevel(logging.ERROR)  # предупреждения о шрифтах — не ошибки, не засоряем вывод
# таблицы в PDF бывают с линиями и без: пробуем оба способа и берём тот, где больше чисел
TEXT_MODE = set()  # страницы, где таблица разобрана «по тексту» (без линий) — числа стоит выборочно сверить
STRATEGIES = [{}, {"vertical_strategy": "text", "horizontal_strategy": "text"}]
NUM = re.compile(r"^\(?[−-]?\d{1,3}(?:[  ]\d{3})*(?:[.,]\d+)?\)?$|^\(?[−-]?\d+(?:[.,]\d+)?\)?$")


GROUPED = re.compile(r"^\(?[−-]?\d{1,3}(?: \d{3})+\)?$")


def merge_groups(row):
    """Разбор таблицы «по тексту» режет число на границе колонок: «1 24» «0 000», «(9 0» «00)», «905» «000».
    Склеиваем соседние ячейки, если вместе они дают число с разрядами через пробел."""
    out = []
    for c in row:
        s = "" if c is None else " ".join(str(c).split())
        prev = out[-1] if out and isinstance(out[-1], str) else ""
        if prev and s and prev[-1].isdigit() and s[0].isdigit():
            last = prev.split(" ")[-1].lstrip("(−-")
            cut = " " in prev and len(last) < 3          # «1 24» — последняя группа разрядов не дописана
            tail = re.match(r"^\d{3}\)?$", s) is not None  # «000» — отдельная группа разрядов, не новое число
            joined = prev + s if cut else prev + " " + s
            if (cut or tail) and GROUPED.match(joined):
                out[-1] = joined
                continue
        out.append(s if c is not None else None)
    return out


def pages_arg(pdf, spec):
    """«1-5» или «3» → индексы страниц (с нуля)."""
    n = len(pdf.pages)
    if not spec:
        return range(n)
    a, _, b = spec.partition("-")
    a, b = int(a), int(b or a)
    return range(max(a, 1) - 1, min(b, n))


def to_num(s):
    """Строка ячейки → число, если это число в принятых в отчётности форматах; иначе — исходная строка."""
    if s is None:
        return None
    t = " ".join(str(s).split())
    if not t or not NUM.match(t):
        return t
    neg = t.startswith("(") or t.lstrip("(").startswith(("-", "−"))
    t = t.strip("()").lstrip("-−").replace(" ", "").replace(" ", "").replace(",", ".")
    try:
        v = float(t)
    except ValueError:
        return s
    v = -v if neg else v
    return int(v) if v.is_integer() else v


def cmd_info(path):
    with pdfplumber.open(path) as pdf:
        n = len(pdf.pages)
        chars = [len(p.chars) for p in pdf.pages[:10]]
    scan = sum(chars) < 50 * len(chars)
    print(f"{Path(path).name}: страниц {n}; " + ("похоже на СКАН — текста нет, прочитать нельзя (нужно распознавание)"
                                               if scan else "текстовый слой есть — можно читать текст и таблицы"))


def cmd_text(path, spec=None):
    src = Path(path)
    out = src.with_suffix(".txt")
    parts = []
    with pdfplumber.open(src) as pdf:
        for i in pages_arg(pdf, spec):
            parts.append(f"\n=== страница {i + 1} ===\n" + (pdf.pages[i].extract_text() or "(нет текста)"))
    text = "".join(parts)
    out.write_text(text, encoding="utf-8")
    print(f"Текст: {out} ({len(text)} знаков)")
    print(text[:1500])


def extract_tables(path, spec=None):
    found = []
    with pdfplumber.open(path) as pdf:
        for i in pages_arg(pdf, spec):
            best, score = [], -1
            for st in STRATEGIES:
                tabs = []
                for t in pdf.pages[i].extract_tables(st):
                    rows = [[to_num(c) for c in merge_groups(r) if c != ""] for r in t if any(c not in (None, "") for c in r)]
                    if rows:
                        tabs.append(rows)
                s = sum(isinstance(c, (int, float)) for rows in tabs for r in rows for c in r)
                if s > score:
                    best, score = tabs, s
                    if st:
                        TEXT_MODE.add(i + 1)
            found += [(i + 1, rows) for rows in best]
    return found


def cmd_tables(path, spec=None):
    src = Path(path)
    found = extract_tables(src, spec)
    if not found:
        sys.exit("Таблиц не нашёл (в PDF они могут быть нарисованы без линий) — используй text и читай цифры из текста.")
    out = src.with_name(src.stem + "_таблицы.xlsx")
    wb = Workbook()
    wb.remove(wb.active)
    for k, (page, rows) in enumerate(found, 1):
        ws = wb.create_sheet(f"с{page}_т{k}"[:31])
        for r in rows:
            ws.append(r)
    wb.save(out)
    print(f"Таблиц: {len(found)} → {out}")
    for page, rows in found[:3]:
        print(f"\n— страница {page}, строк {len(rows)}:")
        for r in rows[:6]:
            print("  " + " | ".join("" if c is None else str(c) for c in r))


def cmd_codes(path):
    src = Path(path)
    found = extract_tables(src)
    if not found:
        sys.exit("Таблиц не нашёл — коды РСБУ искать негде. Попробуй text.")
    tmp = src.with_name(src.stem + "_таблицы.xlsx")
    wb = Workbook()
    wb.remove(wb.active)
    for k, (page, rows) in enumerate(found, 1):
        ws = wb.create_sheet(f"с{page}_т{k}"[:31])
        for r in rows:
            ws.append([str(c).strip() if isinstance(c, str) else c for c in r])
    wb.save(tmp)
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import xl
    xl.cmd_codes(str(tmp))
    if TEXT_MODE:
        print(f"⚠ Таблица без линий (стр. {', '.join(map(str, sorted(TEXT_MODE)))}): сверь 2–3 числа с PDF.")
    print(f"(таблицы сохранены: {tmp})")


if __name__ == "__main__":
    cmds = {"info": cmd_info, "text": cmd_text, "tables": cmd_tables, "codes": cmd_codes}
    if len(sys.argv) < 3 or sys.argv[1] not in cmds:
        print(__doc__)
        sys.exit(1)
    cmds[sys.argv[1]](*sys.argv[2:4])
