#!/usr/bin/env python3
"""Допущения финмодели и пересчёт модели по её собственным формулам (LibreOffice) — для стресс-теста по выбору гостя.

  python3 assume.py list    model.xlsx                        допущения модели: лист «Допущения» (или блок с таким
                                                              заголовком) — номер · подпись · ячейка · значение
  python3 assume.py extract model.xlsx map.csv out.csv        выписать потоки по карте ячеек → out.csv для cf.py
  python3 assume.py scen    model.xlsx map.csv out.csv --set "<ячейки>=<изменение>" [--set …] [--save копия.xlsx]
                                                              изменить допущения, пересчитать модель, выписать потоки

Изменение: «+10%» / «-15%» — относительно значения; «+3пп» — плюс 3 процентных пункта (для долей вида 0,18 — плюс 0,03);
«*1.1» — умножить; «Лист!C5==0.25» — поставить значение. Ячейки: «Допущения!C5», «Модель!B4:F4» (диапазон — каждой ячейке).

Карта map.csv («;»): показатель;ячейки;множитель — например
  год;Модель!B3:F3;            выручка;Модель!B4:F4;1          проценты;Модель!B12:F12;0.001   (тыс. → млн и т. п.)
  деньги_нач;Модель!B16;       прочие_обяз;0                   (0 — статьи нет в модели)
Показатели — столбцы cf.py (выручка, ebitda, налог, capex, изм_оборотки, проценты, погашение, лизинг, прочие_обяз,
деньги_нач; по возможности инвестиции, привлечение, долг, чистая_прибыль, плав_долг).

Исходный файл не меняется: модель открывается в LibreOffice без интерфейса, изменения живут только в памяти
(копия — только с --save). Запуск системным python3 (нужны пакеты libreoffice-calc-nogui и python3-uno).
"""
import csv
import os
import re
import shutil
import socket
import subprocess
import sys
import tempfile
import time
from pathlib import Path

ASSUME_WORDS = ("допущ", "предпосыл", "assumption", "вводн", "input", "параметр", "драйвер")


def split_ref(ref):
    sheet, _, rng = ref.rpartition("!")
    return sheet.strip("'"), rng


def cells_of(rng):
    """«B4:F4» → ['B4', 'C4', …] (по строкам)."""
    from openpyxl.utils import range_boundaries, get_column_letter
    c1, r1, c2, r2 = range_boundaries(rng)
    return [f"{get_column_letter(c)}{r}" for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


# ---------- список допущений (без LibreOffice) ----------

def cmd_list(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    wf = load_workbook(path, data_only=False)
    found = []
    sheets = [ws for ws in wf.worksheets if any(w in ws.title.lower() for w in ASSUME_WORDS)]
    if sheets:
        for ws in sheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                        found.append((ws.title, c, label(ws, c)))
        where = "лист «" + "», «".join(ws.title for ws in sheets) + "»"
    else:
        # блок с заголовком «Допущения» на листе модели: числа-константы правее/ниже заголовка
        for ws in wf.worksheets:
            for row in ws.iter_rows():
                for h in row:
                    if isinstance(h.value, str) and any(w in h.value.lower() for w in ASSUME_WORDS):
                        for r in range(h.row, min(h.row + 40, ws.max_row) + 1):
                            for col in range(h.column, min(h.column + 4, ws.max_column) + 1):
                                c = ws.cell(row=r, column=col)
                                if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                                    found.append((ws.title, c, label(ws, c)))
        where = "блок «Допущения» на листе модели"
    seen, out = set(), []
    for s, c, lab in found:
        if (s, c.coordinate) not in seen:
            seen.add((s, c.coordinate))
            out.append((s, c, lab))
    if not out:
        print("Листа или блока «Допущения» нет. Стресс — стандартными сценариями cf.py (доходы, расходы, ставка, "
              "проценты): python3 cf.py eval … --shock")
        return
    print(f"Допущения модели ({where}), числа-константы:")
    for i, (s, c, lab) in enumerate(out, 1):
        print(f"{i:>3}. {lab or '(без подписи)'} — {s}!{c.coordinate} = {fmt(c.value)}")
    print("Изменить: python3 assume.py scen model.xlsx map.csv out.csv --set \"<лист>!<ячейка>=+10%\"")


def label(ws, c):
    """Подпись ячейки: ближайший текст левее в той же строке, плюс год из заголовка столбца, если он есть."""
    lab = ""
    for col in range(c.column - 1, 0, -1):
        v = ws.cell(row=c.row, column=col).value
        if isinstance(v, str) and v.strip():
            lab = v.strip()
            break
    for r in range(c.row - 1, max(0, c.row - 30), -1):
        v = ws.cell(row=r, column=c.column).value
        if re.fullmatch(r"\s*(19|20)\d\d\D{0,3}\s*", str(v if v is not None else "")):
            return f"{lab} ({str(v).strip()})"
    return lab


def fmt(v):
    return f"{v:,.4g}".replace(",", " ").replace(".", ",") if isinstance(v, float) else str(v)


# ---------- LibreOffice ----------

class Calc:
    """Модель в LibreOffice без интерфейса: читать и менять ячейки, пересчитывать по формулам модели."""

    def __init__(self, path):
        import uno
        from com.sun.star.beans import PropertyValue
        self.uno, self.PV = uno, PropertyValue
        exe = shutil.which("soffice") or shutil.which("libreoffice")
        if not exe:
            sys.exit("Нет LibreOffice: sudo apt install -y libreoffice-calc-nogui python3-uno")
        self.profile = tempfile.mkdtemp(prefix="lo_")
        with socket.socket() as s:
            s.bind(("127.0.0.1", 0))
            port = s.getsockname()[1]
        self.proc = subprocess.Popen(
            [exe, "--headless", "--invisible", "--nologo", "--norestore", "--nodefault",
             f"-env:UserInstallation={uno.systemPathToFileUrl(self.profile)}",
             f"--accept=socket,host=127.0.0.1,port={port};urp;"],
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        ctx = uno.getComponentContext()
        resolver = ctx.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", ctx)
        for _ in range(120):
            try:
                rctx = resolver.resolve(f"uno:socket,host=127.0.0.1,port={port};urp;StarOffice.ComponentContext")
                break
            except Exception:
                time.sleep(0.5)
        else:
            self.close()
            sys.exit("LibreOffice не запустился за 60 секунд")
        self.desktop = rctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", rctx)
        self.doc = self.desktop.loadComponentFromURL(
            uno.systemPathToFileUrl(str(Path(path).resolve())), "_blank", 0, (self._pv("Hidden", True),))
        if self.doc is None:
            self.close()
            sys.exit(f"LibreOffice не открыл файл {path}")
        self.doc.calculateAll()

    def _pv(self, name, value):
        p = self.PV()
        p.Name, p.Value = name, value
        return p

    def cell(self, ref):
        sheet, addr = split_ref(ref)
        if not sheet:
            sys.exit(f"Укажи лист: «Лист!{addr}»")
        return self.doc.Sheets.getByName(sheet).getCellRangeByName(addr)

    def values(self, ref):
        sheet, rng = split_ref(ref)
        out = []
        for a in cells_of(rng):
            c = self.cell(f"{sheet}!{a}")
            out.append(c.getString() if c.getType().value == "TEXT" else c.getValue())
        return out

    def change(self, ref, how):
        sheet, rng = split_ref(ref)
        done = []
        for a in cells_of(rng):
            c = self.cell(f"{sheet}!{a}")
            if c.getFormula().startswith("="):
                print(f"⚠ {sheet}!{a} — формула, а не допущение: меняю всё равно, формула заменится числом")
            old = c.getValue()
            new = apply(old, how)
            c.setValue(new)
            done.append((f"{sheet}!{a}", old, new))
        return done

    def recalc(self):
        self.doc.calculateAll()

    def save(self, out):
        self.doc.storeToURL(self.uno.systemPathToFileUrl(str(Path(out).resolve())),
                            (self._pv("FilterName", "Calc MS Excel 2007 XML"),))

    def close(self):
        try:
            if getattr(self, "doc", None) is not None:
                self.doc.close(True)
            if getattr(self, "desktop", None) is not None:
                self.desktop.terminate()
        except Exception:
            pass
        try:
            self.proc.wait(timeout=20)
        except Exception:
            self.proc.kill()
        shutil.rmtree(self.profile, ignore_errors=True)


def apply(old, how):
    h = how.strip().replace(",", ".").replace(" ", "").lower()
    m = re.fullmatch(r"([+-]\d+(?:\.\d+)?)%", h)
    if m:
        return old * (1 + float(m.group(1)) / 100)
    m = re.fullmatch(r"([+-]\d+(?:\.\d+)?)(пп|п\.п\.|pp)", h)
    if m:
        d = float(m.group(1))
        return old + (d / 100 if abs(old) < 1 else d)
    m = re.fullmatch(r"\*(\d+(?:\.\d+)?)", h)
    if m:
        return old * float(m.group(1))
    m = re.fullmatch(r"=(-?\d+(?:\.\d+)?)", h)
    if m:
        return float(m.group(1))
    sys.exit(f"Не понял изменение «{how}». Примеры: +10%, -15%, +3пп, *1.1, =0.25")


# ---------- карта ячеек → CSV для cf.py ----------

def read_map(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = [r for r in csv.reader(f, delimiter=";") if r and r[0].strip() and not r[0].startswith("#")]
    out = {}
    for r in rows:
        name, ref = r[0].strip(), (r[1].strip() if len(r) > 1 else "")
        k = float((r[2].strip() or "1").replace(",", ".")) if len(r) > 2 and r[2].strip() else 1.0
        out[name] = (ref, k)
    if "год" not in out:
        sys.exit("В карте нет строки «год» (ячейки с годами).")
    return out


def extract(calc, mp, out):
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from cf import COLS, NEED
    years = [str(int(y)) if isinstance(y, float) and y.is_integer() else str(y).strip() for y in calc.values(mp["год"][0])]
    n = len(years)
    table = {"год": years}
    for name, (ref, k) in mp.items():
        if name == "год":
            continue
        if re.fullmatch(r"-?\d+(?:[.,]\d+)?", ref):
            table[name] = [float(ref.replace(",", "."))] * n
            continue
        v = [x * k if isinstance(x, float) else 0.0 for x in calc.values(ref)]
        table[name] = v if name == "деньги_нач" else (v + [0.0] * n)[:n]
    miss = [c for c in NEED if c not in table]
    if miss:
        print("⚠ Нет в карте → 0: " + ", ".join(miss) + " (проверь, что этих статей действительно нет в модели)")
    cols = [c for c in COLS if c in table or c in NEED]
    with open(out, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(cols)
        for i in range(n):
            row = []
            for c in cols:
                if c == "год":
                    row.append(years[i])
                elif c == "деньги_нач":
                    row.append(round(table[c][0], 6) if i == 0 and c in table else "")
                else:
                    row.append(round(table.get(c, [0.0] * n)[i], 6))
            w.writerow(row)
    print(f"Потоки: {out} ({n} лет: {years[0]}–{years[-1]})")


def main():
    a = sys.argv[1:]
    if len(a) >= 2 and a[0] == "list":
        return cmd_list(a[1])
    if len(a) >= 4 and a[0] in ("extract", "scen"):
        model, mp, out = a[1], read_map(a[2]), a[3]
        sets = [a[i + 1] for i, x in enumerate(a) if x == "--set" and i + 1 < len(a)]
        save = a[a.index("--save") + 1] if "--save" in a else None
        if a[0] == "scen" and not sets:
            sys.exit("Нет изменений: --set \"Лист!C5=+10%\"")
        calc = Calc(model)
        try:
            for s in sets:
                ref, _, how = s.partition("=")      # «Лист!C5==0.25» → поставить 0,25
                for cell, old, new in calc.change(ref, how):
                    print(f"  {cell}: {fmt(old)} → {fmt(new)}")
            if sets:
                calc.recalc()
            extract(calc, mp, out)
            if save:
                calc.save(save)
                print(f"Модель со сценарием: {save}")
        finally:
            calc.close()
        return
    print(__doc__)
    sys.exit(1)


if __name__ == "__main__":
    main()
